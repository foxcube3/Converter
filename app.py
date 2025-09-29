import os
import io
import tempfile
from datetime import datetime
from typing import Optional, Tuple

from flask import Flask, render_template, request, send_file, redirect, url_for, flash
from werkzeug.utils import secure_filename

import chardet

# Optional imports - we handle absence gracefully
try:
    from pdfminer.high_level import extract_text as pdf_extract_text
except Exception:
    pdf_extract_text = None

try:
    import docx2txt
except Exception:
    docx2txt = None

try:
    from pptx import Presentation
except Exception:
    Presentation = None

try:
    import openpyxl
except Exception:
    openpyxl = None

try:
    from bs4 import BeautifulSoup
except Exception:
    BeautifulSoup = None

try:
    import ebooklib
    from ebooklib import epub
except Exception:
    ebooklib = None
    epub = None

try:
    from striprtf.striprtf import rtf_to_text
except Exception:
    rtf_to_text = None

try:
    from PIL import Image
except Exception:
    Image = None

try:
    import pytesseract
except Exception:
    pytesseract = None

try:
    import whisper  # openai-whisper
except Exception:
    whisper = None

ALLOWED_EXTENSIONS = None  # accept any; used only for display
UPLOAD_FOLDER = os.path.join(os.path.dirname(__file__), "uploads")
os.makedirs(UPLOAD_FOLDER, exist_ok=True)

app = Flask(__name__)
app.secret_key = os.environ.get("APP_SECRET_KEY", "dev-secret")


def detect_encoding_and_read(bytes_data: bytes) -> str:
    """Attempt to decode arbitrary bytes as text."""
    detection = chardet.detect(bytes_data) or {}
    enc = detection.get("encoding") or "utf-8"
    try:
        return bytes_data.decode(enc, errors="replace")
    except Exception:
        return bytes_data.decode("utf-8", errors="replace")


def extract_text_generic(file_path: str, filename: str) -> Tuple[str, Optional[str]]:
    """
    Extract text from the given file. Returns (text, method_description).
    Best-effort: supports many common formats and falls back to byte decoding.
    """
    ext = os.path.splitext(filename.lower())[1]

    # Plain text and markdown
    if ext in {".txt", ".md", ".log", ".csv", ".json", ".yaml", ".yml"}:
        with open(file_path, "rb") as f:
            content = f.read()
        return detect_encoding_and_read(content), f"Decoded as text ({ext})"

    # PDF
    if ext == ".pdf":
        if pdf_extract_text:
            try:
                text = pdf_extract_text(file_path) or ""
                return text, "pdfminer.six"
            except Exception as e:
                pass
        # Fallback to bytes decode
        with open(file_path, "rb") as f:
            return detect_encoding_and_read(f.read()), "Fallback byte decode (PDF)"

    # DOCX
    if ext == ".docx":
        if docx2txt:
            try:
                text = docx2txt.process(file_path) or ""
                return text, "docx2txt"
            except Exception:
                pass
        with open(file_path, "rb") as f:
            return detect_encoding_and_read(f.read()), "Fallback byte decode (DOCX)"

    # PPTX
    if ext == ".pptx" and Presentation:
        try:
            prs = Presentation(file_path)
            chunks = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        chunks.append(shape.text)
            return "\n\n".join(chunks), "python-pptx"
        except Exception:
            with open(file_path, "rb") as f:
                return detect_encoding_and_read(f.read()), "Fallback byte decode (PPTX)"

    # XLSX
    if ext == ".xlsx" and openpyxl:
        try:
            wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
            chunks = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    row_str = " ".join("" if (cell is None) else str(cell) for cell in row)
                    if row_str.strip():
                        chunks.append(row_str)
            return "\n".join(chunks), "openpyxl"
        except Exception:
            with open(file_path, "rb") as f:
                return detect_encoding_and_read(f.read()), "Fallback byte decode (XLSX)"

    # RTF
    if ext == ".rtf" and rtf_to_text:
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                content = f.read()
            return rtf_to_text(content), "striprtf"
        except Exception:
            with open(file_path, "rb") as f:
                return detect_encoding_and_read(f.read()), "Fallback byte decode (RTF)"

    # HTML
    if ext in {".html", ".htm"} and BeautifulSoup:
        try:
            with open(file_path, "rb") as f:
                raw = f.read()
            soup = BeautifulSoup(raw, "html.parser")
            return soup.get_text(separator="\n"), "beautifulsoup4"
        except Exception:
            with open(file_path, "rb") as f:
                return detect_encoding_and_read(f.read()), "Fallback byte decode (HTML)"

    # EPUB
    if ext == ".epub":
        if epub and BeautifulSoup:
            try:
                book = epub.read_epub(file_path)
                chunks = []
                for item in book.get_items():
                    # Only document items contain XHTML content
                    if hasattr(item, "get_type") and item.get_type() == (getattr(ebooklib, "ITEM_DOCUMENT", None)):
                        content = item.get_content()
                        soup = BeautifulSoup(content, "html.parser")
                        text = soup.get_text(separator="\n")
                        if text.strip():
                            chunks.append(text)
                return "\n\n".join(chunks), "ebooklib + BeautifulSoup"
            except Exception:
                pass
        # Fallback to bytes decode
        with open(file_path, "rb") as f:
            return detect_encoding_and_read(f.read()), "Fallback byte decode (EPUB)"

    # Images - OCR
    if ext in {".png", ".jpg", ".jpeg", ".tiff", ".bmp", ".gif"}:
        if Image and pytesseract:
            try:
                img = Image.open(file_path)
                text = pytesseract.image_to_string(img)
                return text or "", "pytesseract (OCR)"
            except Exception:
                pass
        return "", "OCR not available (requires Tesseract + pytesseract)"

    # Audio/Video - Whisper (best-effort)
    if ext in {".mp3", ".wav", ".m4a", ".mp4", ".webm", ".ogg"} and whisper:
        try:
            # Use a small model for speed; users can change via env
            model_name = os.environ.get("WHISPER_MODEL", "small")
            model = whisper.load_model(model_name)
            result = model.transcribe(file_path)
            return result.get("text", ""), f"openai-whisper ({model_name})"
        except Exception:
            return "", "Audio transcription failed"

    # Fallback: try to decode bytes
    with open(file_path, "rb") as f:
        content = f.read()
    return detect_encoding_and_read(content), "Fallback byte decode"


@app.route("/", methods=["GET", "POST"])
def index():
    if request.method == "POST":
        if "file" not in request.files or request.files["file"].filename == "":
            flash("Please choose a file to upload.")
            return redirect(url_for("index"))
        file = request.files["file"]
        filename = secure_filename(file.filename)

        # Save to temp file
        tmpdir = tempfile.mkdtemp(prefix="upload_")
        saved_path = os.path.join(tmpdir, filename)
        file.save(saved_path)

        text, method = extract_text_generic(saved_path, filename)

        # Create a downloadable text file in memory
        output_txt_name = f"{os.path.splitext(filename)[0]}_text_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.txt"
        buf = io.BytesIO()
        buf.write(text.encode("utf-8", errors="replace"))
        buf.seek(0)

        return send_file(
            buf,
            as_attachment=True,
            download_name=output_txt_name,
            mimetype="text/plain",
            etag=False,
        )

    return render_template("index.html")


@app.route("/health")
def health():
    return {"status": "ok"}


if __name__ == "__main__":
    port = int(os.environ.get("PORT", "5000"))
    app.run(host="0.0.0.0", port=port, debug=True)